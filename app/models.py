import os
import uuid
import json
import datetime
from enum import Enum

class UserRole(Enum):
    ADMIN = 'admin'
    SUPERVISOR = 'supervisor'
    MANAGER = 'manager'

class User:
    """Модель пользователя с ролями"""
    def __init__(self, id=None, username=None, password_hash=None, full_name=None, role='manager', department=None, manager_id=None):
        self.id = id
        self.username = username
        self.password_hash = password_hash
        self.full_name = full_name
        self.role = role
        self.department = department
        self.manager_id = manager_id
        self.is_active = True
        self.created_at = datetime.datetime.now()
        self.last_login = None
    
    def to_dict(self):
        return {
            'id': self.id,
            'username': self.username,
            'full_name': self.full_name,
            'role': self.role,
            'department': self.department,
            'manager_id': self.manager_id,
            'is_active': self.is_active,
            'created_at': self.created_at.isoformat() if self.created_at else None,
            'last_login': self.last_login.isoformat() if self.last_login else None
        }

class TrainingPlan:
    """Модель плана обучения"""
    def __init__(self, id=None, user_id=None, created_at=None, status='active'):
        self.id = id
        self.user_id = user_id
        self.created_at = created_at or datetime.datetime.now()
        self.status = status
        self.weeks = []
        self.total_score = 0
        self.criteria_scores = {}
    
    def generate_from_analysis(self, analysis_result, user_id):
        """Генерация плана обучения на основе анализа"""
        self.user_id = user_id
        self.criteria_scores = analysis_result.get('criteria_scores', {})
        self.total_score = analysis_result.get('total_score', 0)
        
        # Анализ слабых мест
        weak_areas = []
        medium_areas = []
        strong_areas = []
        
        criteria_names = {
            'greeting': 'Приветствие и представление',
            'client_info': 'Сбор информации о клиенте',
            'communication_style': 'Стиль общения',
            'problem_identification': 'Выявление потребности',
            'offers_promotions': 'Рассказ об акциях',
            'solution_proposal': 'Предложение решения',
            'active_listening': 'Активное слушание',
            'closing_efficiency': 'Эффективность завершения'
        }
        
        for crit_id, score_data in self.criteria_scores.items():
            score = score_data.get('score', 0)
            name = criteria_names.get(crit_id, crit_id)
            if score < 40:
                weak_areas.append({'id': crit_id, 'name': name, 'score': score})
            elif score < 70:
                medium_areas.append({'id': crit_id, 'name': name, 'score': score})
            else:
                strong_areas.append({'id': crit_id, 'name': name, 'score': score})
        
        # Формирование недельных планов
        if weak_areas:
            week1 = {
                'week': 1,
                'title': 'Неделя 1: Интенсивный курс (Критические ошибки)',
                'color': '#f44336',
                'topics': [
                    {
                        'criterion': item['name'],
                        'current_score': item['score'],
                        'target_score': 70,
                        'exercises': self._get_exercises_for_criterion(item['id'])
                    } for item in weak_areas
                ],
                'total_hours': len(weak_areas) * 3,
                'status': 'pending'
            }
            self.weeks.append(week1)
        
        if medium_areas:
            week2 = {
                'week': 2,
                'title': 'Неделя 2: Закрепление и развитие навыков',
                'color': '#ff9800',
                'topics': [
                    {
                        'criterion': item['name'],
                        'current_score': item['score'],
                        'target_score': 85,
                        'exercises': self._get_exercises_for_criterion(item['id'])
                    } for item in medium_areas
                ],
                'total_hours': len(medium_areas) * 2,
                'status': 'pending'
            }
            self.weeks.append(week2)
        
        if strong_areas or not (weak_areas or medium_areas):
            week3 = {
                'week': 3,
                'title': 'Неделя 3: Продвинутый уровень и менторство',
                'color': '#4caf50',
                'topics': [
                    {
                        'criterion': item['name'],
                        'current_score': item['score'],
                        'target_score': 95,
                        'exercises': [
                            'Участие в ролевых играх как наставник',
                            'Разбор сложных кейсов с новыми сотрудниками',
                            'Сертификация и тестирование навыков'
                        ]
                    } for item in (strong_areas if strong_areas else [{'id': 'general', 'name': 'Поддержание высокого уровня', 'score': self.total_score}])
                ],
                'total_hours': 8,
                'status': 'pending'
            }
            self.weeks.append(week3)
        
        return self
    
    def _get_exercises_for_criterion(self, criterion_id):
        """Получение упражнений для конкретного критерия"""
        exercises_map = {
            'greeting': [
                'Прослушивание эталонных приветствий',
                'Отработка стандартного скрипта приветствия',
                'Ролевая игра "Первый контакт с клиентом"'
            ],
            'client_info': [
                'Составление чек-листа сбора информации',
                'Практика активных вопросов',
                'Анализ успешных диалогов коллег'
            ],
            'communication_style': [
                'Тренинг по вежливому общению',
                'Анализ тональности своих разговоров',
                'Практика позитивных формулировок'
            ],
            'problem_identification': [
                'Изучение методики SPIN-продаж',
                'Практика выявления скрытых потребностей',
                'Разбор кейсов с возражениями'
            ],
            'offers_promotions': [
                'Изучение актуального прайс-листа',
                'Отработка презентации спецпредложений',
                'Ролевая игра "Акции и скидки"'
            ],
            'solution_proposal': [
                'Составление шаблонов предложений',
                'Практика завершения сделки',
                'Анализ конверсии своих предложений'
            ],
            'active_listening': [
                'Упражнения на перефразирование',
                'Практика эмпатического слушания',
                'Тренинг "Слышу и понимаю"'
            ],
            'closing_efficiency': [
                'Отработка стандартных фраз завершения',
                'Практика договоренностей о следующих шагах',
                'Анализ завершенных диалогов'
            ]
        }
        return exercises_map.get(criterion_id, ['Индивидуальная работа с наставником', 'Повторение стандартов качества'])
    
    def to_dict(self):
        return {
            'id': self.id,
            'user_id': self.user_id,
            'created_at': self.created_at.isoformat() if self.created_at else None,
            'status': self.status,
            'total_score': self.total_score,
            'criteria_scores': self.criteria_scores,
            'weeks': self.weeks
        }

class ReportGenerator:
    """Генератор отчётов"""
    
    @staticmethod
    def generate_html_report(analysis_result, user_info=None, training_plan=None):
        """Генерация HTML отчёта"""
        html = f"""
        <!DOCTYPE html>
        <html lang="ru">
        <head>
            <meta charset="UTF-8">
            <title>Отчёт об анализе разговора - COFRESH Voice Analytics</title>
            <style>
                body {{ font-family: 'Segoe UI', Arial, sans-serif; margin: 0; padding: 20px; background: #f5f5f5; }}
                .container {{ max-width: 1200px; margin: 0 auto; background: white; border-radius: 15px; box-shadow: 0 5px 20px rgba(0,0,0,0.1); overflow: hidden; }}
                .header {{ background: linear-gradient(135deg, #0a2a4a, #2a7bb0); color: white; padding: 30px; text-align: center; }}
                .header h1 {{ margin: 0; font-size: 28px; }}
                .header p {{ margin: 10px 0 0; opacity: 0.9; }}
                .content {{ padding: 30px; }}
                .score-section {{ text-align: center; padding: 20px; background: #f8f9fa; border-radius: 15px; margin-bottom: 30px; }}
                .score {{ font-size: 64px; font-weight: bold; color: #ff9800; }}
                .grade {{ display: inline-block; padding: 8px 24px; border-radius: 40px; font-weight: bold; margin-top: 10px; }}
                .grade-excellent {{ background: #2e7d32; color: white; }}
                .grade-good {{ background: #2a7bb0; color: white; }}
                .grade-average {{ background: #4a90b0; color: white; }}
                .criteria-grid {{ display: grid; grid-template-columns: repeat(2, 1fr); gap: 20px; margin-bottom: 30px; }}
                .criteria-item {{ background: #f8f9fa; border-radius: 12px; padding: 15px; }}
                .criteria-name {{ font-weight: bold; margin-bottom: 10px; }}
                .criteria-score {{ font-size: 24px; font-weight: bold; color: #ff9800; }}
                .progress-bar {{ height: 8px; background: #e0e0e0; border-radius: 4px; overflow: hidden; margin: 10px 0; }}
                .progress-fill {{ height: 100%; border-radius: 4px; transition: width 0.3s; }}
                .transcript {{ background: #f8f9fa; border-radius: 12px; padding: 20px; margin-top: 30px; }}
                .transcript-line {{ padding: 12px; border-left: 4px solid; margin-bottom: 10px; background: white; border-radius: 8px; }}
                .transcript-line.admin {{ border-left-color: #2a7bb0; }}
                .transcript-line.client {{ border-left-color: #4a90b0; }}
                .training-plan {{ background: #e8f5e9; border-radius: 12px; padding: 20px; margin-top: 30px; }}
                .week {{ background: white; border-radius: 10px; padding: 15px; margin-bottom: 15px; }}
                .week-title {{ font-weight: bold; font-size: 18px; margin-bottom: 10px; }}
                .footer {{ text-align: center; padding: 20px; background: #f8f9fa; color: #666; font-size: 12px; }}
                table {{ width: 100%; border-collapse: collapse; }}
                th, td {{ padding: 10px; text-align: left; border-bottom: 1px solid #ddd; }}
                @media (max-width: 768px) {{ .criteria-grid {{ grid-template-columns: 1fr; }} }}
            </style>
        </head>
        <body>
            <div class="container">
                <div class="header">
                    <h1>COFRESH Voice Analytics Pro</h1>
                    <p>Отчёт об анализе качества телефонного разговора</p>
                </div>
                <div class="content">
                    <div class="score-section">
                        <div class="score">{analysis_result.get('total_score', 0)}<span style="font-size: 24px;">/100</span></div>
                        <div class="grade grade-{analysis_result.get('grade_class', '')}">{analysis_result.get('grade', '')}</div>
                        <p>{analysis_result.get('grade_description', '')}</p>
                    </div>
                    
                    <h3>📊 Оценка по критериям</h3>
                    <div class="criteria-grid">
        """
        
        criteria_names = {
            'greeting': '👋 Приветствие и представление',
            'client_info': '📝 Сбор информации',
            'communication_style': '💬 Стиль общения',
            'problem_identification': '🎯 Выявление потребности',
            'offers_promotions': '🏷️ Рассказ об акциях',
            'solution_proposal': '💡 Предложение решения',
            'active_listening': '👂 Активное слушание',
            'closing_efficiency': '✅ Завершение разговора'
        }
        
        for crit_id, name in criteria_names.items():
            score_data = analysis_result.get('criteria_scores', {}).get(crit_id, {'score': 0})
            score = score_data.get('score', 0)
            percentage = (score / 100) * 100
            bar_color = '#2e7d32' if percentage >= 80 else '#ff9800' if percentage >= 60 else '#f44336'
            
            html += f"""
                        <div class="criteria-item">
                            <div class="criteria-name">{name}</div>
                            <div class="criteria-score">{score}/100</div>
                            <div class="progress-bar">
                                <div class="progress-fill" style="width: {percentage}%; background: {bar_color};"></div>
                            </div>
                            <div style="font-size: 12px; color: #666; margin-top: 8px;">
                                Найдено ключевых слов: {score_data.get('keyword_count', 0)}
                            </div>
                        </div>
            """
        
        html += """
                    </div>
                    
                    <h3>📈 Статистика разговора</h3>
                    <table>
        """
        
        html += f"""
                        <tr><th>Параметр</th><th>Значение</th></tr>
                        <tr><td>Длительность</td><td>{analysis_result.get('duration_formatted', '0:00')}</td></tr>
                        <tr><td>Всего слов</td><td>{analysis_result.get('word_count', 0)}</td></tr>
                        <tr><td>Слов администратора</td><td>{analysis_result.get('admin_word_count', 0)} ({analysis_result.get('dialogue_analysis', {}).get('admin_percentage', 0)}%)</td></tr>
                        <tr><td>Слов клиента</td><td>{analysis_result.get('client_word_count', 0)} ({analysis_result.get('dialogue_analysis', {}).get('client_percentage', 0)}%)</td></tr>
                        <tr><td>Количество вопросов</td><td>{analysis_result.get('dialogue_analysis', {}).get('question_count', 0)}</td></tr>
                        <tr><td>Тональность</td><td>{analysis_result.get('sentiment_label', 'Нейтральная')} (балл: {analysis_result.get('sentiment', 50)})</td></tr>
                    </table>
        """
        
        if training_plan:
            html += f"""
                    <div class="training-plan">
                        <h3>📚 План обучения (персональный)</h3>
                        <p>Итоговый балл: {training_plan.get('total_score', 0)}/100</p>
            """
            for week in training_plan.get('weeks', []):
                html += f"""
                        <div class="week">
                            <div class="week-title" style="color: {week.get('color', '#333')};">📅 {week.get('title', '')}</div>
                            <p><strong>Темы для изучения:</strong></p>
                            <ul>
                """
                for topic in week.get('topics', []):
                    html += f"<li><strong>{topic.get('criterion', '')}</strong> (текущий балл: {topic.get('current_score', 0)}/100 → цель: {topic.get('target_score', 0)}/100)"
                    if topic.get('exercises'):
                        html += "<ul>"
                        for ex in topic.get('exercises', []):
                            html += f"<li>{ex}</li>"
                        html += "</ul>"
                    html += "</li>"
                html += f"""
                            </ul>
                            <p>⏱️ Рекомендуемое время: {week.get('total_hours', 0)} часов</p>
                        </div>
                """
            html += "</div>"
        
        html += f"""
                    <div class="transcript">
                        <h3>💬 Расшифровка диалога</h3>
        """
        
        for segment in analysis_result.get('segments', []):
            speaker_name = 'Администратор' if segment.get('speaker') == 'admin' else 'Клиент'
            speaker_icon = '👨‍💼' if segment.get('speaker') == 'admin' else '👤'
            html += f"""
                        <div class="transcript-line {segment.get('speaker', 'client')}">
                            <strong>{speaker_icon} {speaker_name}</strong> <span style="color: #999;">⏱️ {segment.get('timestamp', '00:00')}</span>
                            <div style="margin-top: 8px;">{segment.get('text', '')}</div>
                        </div>
            """
        
        html += f"""
                    </div>
                </div>
                <div class="footer">
                    <p>COFRESH Voice Analytics Pro · Отчёт сгенерирован {datetime.datetime.now().strftime('%d.%m.%Y %H:%M')}</p>
                    <p>АВТО С ЗАБОТОЙ · ТОП-3 дилер РФ</p>
                </div>
            </div>
        </body>
        </html>
        """
        
        return html
    
    @staticmethod
    def generate_pdf_report(analysis_result, user_info=None, training_plan=None):
        """Генерация PDF отчёта (использует HTML в промежуточном формате)"""
        html = ReportGenerator.generate_html_report(analysis_result, user_info, training_plan)
        import io
        import base64
        from weasyprint import HTML
        
        pdf_file = io.BytesIO()
        HTML(string=html).write_pdf(pdf_file)
        pdf_file.seek(0)
        return pdf_file.getvalue()
    
    @staticmethod
    def generate_excel_report(analysis_result, training_plan=None):
        """Генерация Excel отчёта"""
        import io
        import pandas as pd
        from openpyxl.styles import Font, PatternFill, Alignment
        
        output = io.BytesIO()
        
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            # Лист 1: Общая информация
            summary_data = {
                'Показатель': ['Дата анализа', 'Длительность', 'Общий балл', 'Оценка', 'Всего слов', 
                              'Слов администратора', 'Слов клиента', 'Тональность', 'Точность распознавания'],
                'Значение': [datetime.datetime.now().strftime('%d.%m.%Y %H:%M'),
                           analysis_result.get('duration_formatted', '0:00'),
                           analysis_result.get('total_score', 0),
                           analysis_result.get('grade', ''),
                           analysis_result.get('word_count', 0),
                           analysis_result.get('admin_word_count', 0),
                           analysis_result.get('client_word_count', 0),
                           analysis_result.get('sentiment_label', 'Нейтральная'),
                           analysis_result.get('confidence_percent', '0%')]
            }
            df_summary = pd.DataFrame(summary_data)
            df_summary.to_excel(writer, sheet_name='Общая информация', index=False)
            
            # Лист 2: Оценки критериев
            criteria_data = []
            criteria_names = {
                'greeting': 'Приветствие и представление',
                'client_info': 'Сбор информации',
                'communication_style': 'Стиль общения',
                'problem_identification': 'Выявление потребности',
                'offers_promotions': 'Рассказ об акциях',
                'solution_proposal': 'Предложение решения',
                'active_listening': 'Активное слушание',
                'closing_efficiency': 'Завершение разговора'
            }
            for crit_id, name in criteria_names.items():
                score_data = analysis_result.get('criteria_scores', {}).get(crit_id, {'score': 0, 'keyword_count': 0})
                criteria_data.append({
                    'Критерий': name,
                    'Балл': score_data.get('score', 0),
                    'Найдено ключевых слов': score_data.get('keyword_count', 0),
                    'Ключевые слова': ', '.join(score_data.get('found_keywords', [])) if score_data.get('found_keywords') else '-'
                })
            df_criteria = pd.DataFrame(criteria_data)
            df_criteria.to_excel(writer, sheet_name='Оценки критериев', index=False)
            
            # Лист 3: План обучения (если есть)
            if training_plan:
                training_data = []
                for week in training_plan.get('weeks', []):
                    for topic in week.get('topics', []):
                        training_data.append({
                            'Неделя': week.get('week', 0),
                            'Тема': topic.get('criterion', ''),
                            'Текущий балл': topic.get('current_score', 0),
                            'Целевой балл': topic.get('target_score', 0),
                            'Упражнения': ', '.join(topic.get('exercises', []))
                        })
                df_training = pd.DataFrame(training_data)
                df_training.to_excel(writer, sheet_name='План обучения', index=False)
            
            # Лист 4: Расшифровка диалога
            transcript_data = []
            for idx, segment in enumerate(analysis_result.get('segments', [])):
                transcript_data.append({
                    '№': idx + 1,
                    'Говорящий': 'Администратор' if segment.get('speaker') == 'admin' else 'Клиент',
                    'Время': segment.get('timestamp', '00:00'),
                    'Текст': segment.get('text', ''),
                    'Точность': f"{int(segment.get('confidence', 0.8) * 100)}%"
                })
            df_transcript = pd.DataFrame(transcript_data)
            df_transcript.to_excel(writer, sheet_name='Расшифровка диалога', index=False)
        
        output.seek(0)
        return output.getvalue()
    
    @staticmethod
    def generate_word_report(analysis_result, user_info=None, training_plan=None):
        """Генерация Word отчёта"""
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        
        doc = Document()
        
        # Заголовок
        title = doc.add_heading('COFRESH Voice Analytics Pro', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_heading('Отчёт об анализе качества телефонного разговора', level=1)
        
        # Общая информация
        doc.add_heading('1. Общая информация', level=2)
        info_table = doc.add_table(rows=8, cols=2)
        info_table.style = 'Table Grid'
        info_data = [
            ('Дата анализа', datetime.datetime.now().strftime('%d.%m.%Y %H:%M')),
            ('Длительность', analysis_result.get('duration_formatted', '0:00')),
            ('Общий балл', f"{analysis_result.get('total_score', 0)}/100"),
            ('Оценка', analysis_result.get('grade', '')),
            ('Всего слов', str(analysis_result.get('word_count', 0))),
            ('Слов администратора', str(analysis_result.get('admin_word_count', 0))),
            ('Слов клиента', str(analysis_result.get('client_word_count', 0))),
            ('Тональность', analysis_result.get('sentiment_label', 'Нейтральная'))
        ]
        for i, (key, value) in enumerate(info_data):
            info_table.cell(i, 0).text = key
            info_table.cell(i, 1).text = value
        
        # Оценки критериев
        doc.add_heading('2. Оценка по критериям', level=2)
        criteria_names = {
            'greeting': 'Приветствие и представление',
            'client_info': 'Сбор информации',
            'communication_style': 'Стиль общения',
            'problem_identification': 'Выявление потребности',
            'offers_promotions': 'Рассказ об акциях',
            'solution_proposal': 'Предложение решения',
            'active_listening': 'Активное слушание',
            'closing_efficiency': 'Завершение разговора'
        }
        
        criteria_table = doc.add_table(rows=len(criteria_names) + 1, cols=3)
        criteria_table.style = 'Table Grid'
        headers = criteria_table.rows[0].cells
        headers[0].text = 'Критерий'
        headers[1].text = 'Балл'
        headers[2].text = 'Ключевые слова'
        
        for i, (crit_id, name) in enumerate(criteria_names.items()):
            score_data = analysis_result.get('criteria_scores', {}).get(crit_id, {'score': 0, 'found_keywords': []})
            row = criteria_table.rows[i + 1].cells
            row[0].text = name
            row[1].text = f"{score_data.get('score', 0)}/100"
            row[2].text = ', '.join(score_data.get('found_keywords', [])) if score_data.get('found_keywords') else '-'
        
        # План обучения
        if training_plan:
            doc.add_heading('3. Персональный план обучения', level=2)
            for week in training_plan.get('weeks', []):
                doc.add_heading(week.get('title', ''), level=3)
                for topic in week.get('topics', []):
                    p = doc.add_paragraph()
                    p.add_run(f"• {topic.get('criterion', '')}").bold = True
                    p.add_run(f" (текущий балл: {topic.get('current_score', 0)}/100 → цель: {topic.get('target_score', 0)}/100)")
                    if topic.get('exercises'):
                        doc.add_paragraph('Упражнения:', style='List Bullet')
                        for ex in topic.get('exercises', []):
                            doc.add_paragraph(ex, style='List Bullet 2')
        
        # Расшифровка
        doc.add_heading('4. Расшифровка диалога', level=2)
        for segment in analysis_result.get('segments', []):
            speaker = 'Администратор' if segment.get('speaker') == 'admin' else 'Клиент'
            p = doc.add_paragraph()
            p.add_run(f"[{segment.get('timestamp', '00:00')}] {speaker}: ").bold = True
            p.add_run(segment.get('text', ''))
        
        # Сохранение в байты
        import io
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
